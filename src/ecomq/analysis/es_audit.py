"""E↔S 混淆的标注审计（DEVLOG S17）：抽取系统性错误做盲评，估计标注本身有争议的比例。

sample  从测试集中抽样，写出盲评表（不含标注与模型预测，顺序打乱）和单独的密钥文件
  错误组：标注为 exact 而 ≥k 个模型判为 substitute（E2S），以及反过来（S2E）；每个 query 最多取 1 条
  对照组：所有模型都判对的 exact / substitute，用来衡量判断人自身与标注的一致程度
sheet   把 blind.csv 整理成带下拉菜单和说明的 Excel 盲评表，供人工独立判断（S17 的表是手工整理的，S24 起由此生成）
score   读入判断结果（item -> E/S/C/I/U，U 表示不确定），对照密钥计算各组一致率与“超出对照组的不一致率”

--mode ci（DEVLOG S24）：审计 complement / irrelevant 的标注。按粗粒度类别 R（exact 或 substitute）/ C / I 比较：
  C2R  标注 complement，≥k 个模型判为 R；  I2R  标注 irrelevant，≥k 个判为 R；
  R2I  标注 R，≥k 个判为 irrelevant；      CI   C 与 I 之间互判（标注 C 而 ≥k 个判 I，或反过来）；
  对照组 control_C、control_I：所有模型都判对。评分时判断人的 E 与 S 都算作 R。
"""
import argparse
import csv
import json
import random
from collections import Counter, defaultdict
from pathlib import Path

from ecomq.data.io import read_jsonl
from ecomq.schema import parse_output

ABBR = {"exact": "E", "substitute": "S", "complement": "C", "irrelevant": "I"}
COARSE = {"exact": "R", "substitute": "R", "complement": "C", "irrelevant": "I"}
COARSE_ABBR = {"E": "R", "S": "R", "C": "C", "I": "I", "U": "U"}


def load_votes(pred_paths: list[str]) -> dict[str, list[str]]:
    votes = defaultdict(list)
    for path in pred_paths:
        for p in read_jsonl(path):
            obj = parse_output(p["output"], "relevance") or {}
            votes[p["id"]].append(obj.get("relevance", "invalid"))
    return votes


def select(recs, votes, k: int, n_err: int, n_ctrl: int, seed: int):
    rng = random.Random(seed)
    pools = defaultdict(list)
    for r in recs:
        gold, v = r["label"]["relevance"], votes[r["id"]]
        if gold == "exact" and v.count("substitute") >= k:
            pools["E2S"].append(r)
        elif gold == "substitute" and v.count("exact") >= k:
            pools["S2E"].append(r)
        elif gold in ("exact", "substitute") and all(x == gold for x in v):
            pools["control_" + ABBR[gold]].append(r)

    def pick(pool, n):
        rng.shuffle(pool)
        seen, out = set(), []
        for r in pool:  # 每个 query 最多 1 条，增加多样性
            if r["group"] not in seen:
                seen.add(r["group"])
                out.append(r)
            if len(out) == n:
                break
        return out

    chosen = {"E2S": pick(pools["E2S"], n_err), "S2E": pick(pools["S2E"], n_err),
              "control_E": pick(pools["control_E"], n_ctrl), "control_S": pick(pools["control_S"], n_ctrl)}
    sizes = {name: len(p) for name, p in pools.items()}
    items = [(name, r) for name, rs in chosen.items() for r in rs]
    rng.shuffle(items)
    return items, sizes


def _pick(rng, pool, n):
    rng.shuffle(pool)
    seen, out = set(), []
    for r in pool:  # 每个 query 最多 1 条，增加多样性
        if r["group"] not in seen:
            seen.add(r["group"])
            out.append(r)
        if len(out) == n:
            break
    return out


def select_ci(recs, votes, k: int, n_err: int, n_ci: int, n_ctrl: int, seed: int):
    rng = random.Random(seed)
    pools = defaultdict(list)
    for r in recs:
        gold, v = COARSE[r["label"]["relevance"]], [COARSE.get(x, "?") for x in votes[r["id"]]]
        if gold == "C" and v.count("R") >= k:
            pools["C2R"].append(r)
        elif gold == "I" and v.count("R") >= k:
            pools["I2R"].append(r)
        elif gold == "R" and v.count("I") >= k:
            pools["R2I"].append(r)
        elif (gold == "C" and v.count("I") >= k) or (gold == "I" and v.count("C") >= k):
            pools["CI"].append(r)
        elif gold in ("C", "I") and all(x == gold for x in v):
            pools["control_" + gold].append(r)
    sizes_ = {"C2R": n_err, "I2R": n_err, "R2I": n_err, "CI": n_ci, "control_C": n_ctrl, "control_I": n_ctrl}
    chosen = {name: _pick(rng, pools[name], n) for name, n in sizes_.items()}
    items = [(name, r) for name, rs in chosen.items() for r in rs]
    rng.shuffle(items)
    return items, {name: len(p) for name, p in pools.items()}


def score_ci(key: dict, labels: dict[str, str]) -> dict:
    """粗粒度（R/C/I）比较：判断人的 E、S 都算 R。"""
    by_set = defaultdict(Counter)
    for item, info in key["items"].items():
        lab = labels.get(item)
        if lab is None:
            continue
        lab, gold, model = COARSE_ABBR[lab], info["gold_coarse"], info["model_majority"]
        verdict = "uncertain" if lab == "U" else "agree_gold" if lab == gold else "agree_model" if lab == model else "other"
        by_set[info["set"]][verdict] += 1
        by_set[info["set"]]["n"] += 1
    rate = lambda c, k: round(c[k] / c["n"], 4) if c["n"] else None  # noqa: E731
    out = {name: {**dict(c), **{f"{k}_rate": rate(c, k) for k in ("agree_gold", "agree_model", "other", "uncertain")}}
           for name, c in sorted(by_set.items())}
    ctrl = by_set["control_C"] + by_set["control_I"]
    if ctrl["n"]:
        ctrl_dis = 1 - ctrl["agree_gold"] / ctrl["n"]
        out["summary"] = {"control_disagree_with_gold": round(ctrl_dis, 4)}
        for name in ("C2R", "I2R", "R2I", "CI"):
            c = by_set[name]
            if c["n"]:
                out["summary"][name] = {"disagree_with_gold": round(1 - c["agree_gold"] / c["n"], 4),
                                        "excess_disagreement": round(1 - c["agree_gold"] / c["n"] - ctrl_dis, 4),
                                        "agree_with_model": round(c["agree_model"] / c["n"], 4)}
    return out


def score(key: dict, labels: dict[str, str]) -> dict:
    if key.get("mode") == "ci":
        return score_ci(key, labels)
    by_set = defaultdict(Counter)
    for item, info in key["items"].items():
        lab = labels.get(item)
        if lab is None:
            continue
        gold, model = info["gold"], info["model_majority"]
        if lab == "U":
            by_set[info["set"]]["uncertain"] += 1
        elif lab == gold:
            by_set[info["set"]]["agree_gold"] += 1
        elif lab == model:
            by_set[info["set"]]["agree_model"] += 1
        else:
            by_set[info["set"]]["other"] += 1
        by_set[info["set"]]["n"] += 1

    def rate(c, key_):
        return round(c[key_] / c["n"], 4) if c["n"] else None

    out = {name: {**dict(c), **{f"{k}_rate": rate(c, k) for k in ("agree_gold", "agree_model", "other", "uncertain")}}
           for name, c in sorted(by_set.items())}
    ctrl = by_set["control_E"] + by_set["control_S"]
    err = by_set["E2S"] + by_set["S2E"]
    if ctrl["n"] and err["n"]:
        ctrl_dis = 1 - ctrl["agree_gold"] / ctrl["n"]
        err_dis = 1 - err["agree_gold"] / err["n"]
        out["summary"] = {
            "error_items": err["n"], "control_items": ctrl["n"],
            "error_disagree_with_gold": round(err_dis, 4),
            "control_disagree_with_gold": round(ctrl_dis, 4),
            "excess_disagreement": round(err_dis - ctrl_dis, 4),
            "error_agree_with_model": round(err["agree_model"] / err["n"], 4),
            "error_uncertain": round(err["uncertain"] / err["n"], 4),
        }
    return out


def build_sheet(blind_csv: str, out_xlsx: str, title: str, key_path: str) -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = list(csv.DictReader(open(blind_csv, encoding="utf-8")))
    n = len(rows)
    wb = openpyxl.Workbook()
    intro = wb.active
    intro.title = "说明"
    lines = [
        title, None,
        "怎么做：在「标注」工作表里，逐行阅读 query 和商品信息，在黄色的「你的判断」列中用下拉菜单选择一个标签，可在「备注」列写理由。",
        "请独立判断：表中不包含原始标注和模型预测，顺序已随机打乱。只需要修改黄色列和备注列，其余列不要改。", None,
        "标签定义（与 ESCI 官方定义一致）",
        "E = exact：商品满足 query 的全部要求。",
        "S = substitute：商品缺少 query 的某个要求，但可以作为功能上的替代品。",
        "C = complement：商品不满足 query，但可以与满足 query 的商品搭配使用。",
        "I = irrelevant：商品无关，或不满足 query 的核心要求。",
        "U = 不确定：信息不足以判断。", None,
        f"完成后保存文件，用 python -m ecomq.analysis.es_audit score --key {key_path} --labels <本文件> 计算结果。",
    ]
    for line in lines:
        intro.append([line])
    intro["A1"].font = Font(bold=True, size=13)
    intro.column_dimensions["A"].width = 110
    ws = wb.create_sheet("标注")
    ws.append(["序号", "query", "商品标题", "品牌", "颜色", "商品要点", "你的判断", "备注"])
    for r in rows:
        ws.append([int(r["item"]), r["query"], r["title"], r["brand"], r["color"], r["details"], None, None])
    for col, width in zip("ABCDEFGH", (6, 28, 48, 16, 14, 60, 11, 30)):
        ws.column_dimensions[col].width = width
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for i in range(2, n + 2):
        ws[f"G{i}"].fill = yellow
        for col in "BCF":
            ws[f"{col}{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    dv = DataValidation(type="list", formula1='"E,S,C,I,U"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{n + 1}")
    ws.freeze_panes = "B2"
    summ = wb.create_sheet("汇总")
    summ.append(["判断进度（自动统计，不含答案）", None])
    summ.append(["标签", "条数"])
    for k, lab in enumerate("ESCIU", start=3):
        summ.append([lab, f"=COUNTIF(标注!$G$2:$G${n + 1},A{k})"])
    summ.append(["已判断", "=SUM(B3:B7)"])
    summ.append(["总条数", f"=COUNTA(标注!$A$2:$A${n + 1})"])
    summ.append(["完成率", "=IF(B9=0,0,B8/B9)"])
    wb.save(out_xlsx)
    return n


def read_labels(path: str) -> dict[str, str]:
    """支持 CSV（列 item,label）或 Excel 盲评表（“标注”工作表的“序号”与“你的判断”列）。"""
    norm = {"E": "E", "S": "S", "C": "C", "I": "I", "U": "U", "?": "U", "不确定": "U"}
    if path.endswith(".xlsx"):
        import openpyxl

        ws = openpyxl.load_workbook(path, data_only=True)["标注"]
        header = [c.value for c in ws[1]]
        i_item, i_lab = header.index("序号"), header.index("你的判断")
        rows = ((r[i_item], r[i_lab]) for r in ws.iter_rows(min_row=2, values_only=True))
    else:
        rows = ((r["item"], r["label"]) for r in csv.DictReader(open(path, encoding="utf-8")))
    out = {}
    for item, lab in rows:
        lab = str(lab).strip().upper() if lab is not None else ""
        lab = norm.get(lab, norm.get(lab[:1], None)) if lab else None
        if item is not None and lab:
            out[str(int(item))] = lab
    return out


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample")
    s.add_argument("--test", required=True)
    s.add_argument("--preds", nargs="+", required=True)
    s.add_argument("--k", type=int, default=5)
    s.add_argument("--n-err", type=int, default=50)
    s.add_argument("--n-ctrl", type=int, default=10)
    s.add_argument("--mode", choices=["es", "ci"], default="es")
    s.add_argument("--n-ci", type=int, default=20, help="--mode ci 时 C↔I 互判组的条数")
    s.add_argument("--seed", type=int, default=0)
    s.add_argument("--out-dir", required=True)
    h = sub.add_parser("sheet")
    h.add_argument("--blind", required=True)
    h.add_argument("--key", required=True, help="只写进说明里的评分命令，不读取")
    h.add_argument("--title", required=True)
    h.add_argument("--output", required=True)
    c = sub.add_parser("score")
    c.add_argument("--key", required=True)
    c.add_argument("--labels", required=True)
    c.add_argument("--output")
    args = ap.parse_args()

    if args.cmd == "sample":
        recs = list(read_jsonl(args.test))
        votes = load_votes(args.preds)
        if args.mode == "ci":
            items, sizes = select_ci(recs, votes, args.k, args.n_err, args.n_ci, args.n_ctrl, args.seed)
        else:
            items, sizes = select(recs, votes, args.k, args.n_err, args.n_ctrl, args.seed)
        out = Path(args.out_dir)
        out.mkdir(parents=True, exist_ok=True)
        key = {"mode": args.mode, "k": args.k, "n_models": len(args.preds), "pool_sizes": sizes, "items": {}}
        with open(out / "blind.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["item", "query", "title", "brand", "color", "details"])
            for i, (name, r) in enumerate(items, 1):
                p = r.get("product", {})
                w.writerow([i, r["query"], r["title"], p.get("brand", ""), p.get("color", ""), p.get("bullets", "")])
                v = votes[r["id"]]
                majority = (Counter(COARSE.get(x, "?") for x in v) if args.mode == "ci"
                            else Counter(ABBR.get(x, "?") for x in v)).most_common(1)[0][0]
                key["items"][str(i)] = {"id": r["id"], "set": name, "gold": ABBR[r["label"]["relevance"]],
                                        "gold_coarse": COARSE[r["label"]["relevance"]], "model_majority": majority,
                                        "votes": dict(Counter(ABBR.get(x, x) for x in v)),
                                        "query_id": r["meta"]["query_id"], "product_id": r["meta"]["product_id"]}
        (out / "key.json").write_text(json.dumps(key, indent=1, ensure_ascii=False), encoding="utf-8")
        print(json.dumps({"pool_sizes": sizes, "sampled": Counter(n for n, _ in items)}, ensure_ascii=False))
    elif args.cmd == "sheet":
        print(f"写出 {build_sheet(args.blind, args.output, args.title, args.key)} 条 -> {args.output}")
    else:
        report = score(json.loads(Path(args.key).read_text(encoding="utf-8")), read_labels(args.labels))
        text = json.dumps(report, indent=2, ensure_ascii=False)
        print(text)
        if args.output:
            Path(args.output).write_text(text, encoding="utf-8")


if __name__ == "__main__":
    main()
